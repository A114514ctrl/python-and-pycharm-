# -*- coding: utf-8 -*-
"""
06_make_report.py  生成 Word 实验报告（严格按《课程期末考核要求》的格式与章节）

格式要求（考核要求文档）：
  * 标题：方正小标宋简体 小二（写明具体题目）
  * 姓名    学号
  * 正文：仿宋_GB2312
  * 章节：一、题目介绍 → 二、所用工具介绍 → 三、数据集准备 → 四、完整训练和测试过程
          → 五、方法说明（模型介绍/改进展示/原理） → 六、效果展示 → 七、个人总结心得
  * 子级：（一）（二）…… → 1. 2. …… → （1）（2）……

运行（Python 3.10）：python 06_make_report.py
输出：report/实验报告.docx
"""

import json
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(HERE)
RESULT_DIR = os.path.join(PROJECT, "results")
SHOT_DIR = os.path.join(RESULT_DIR, "console_shots")
CODE_SHOT_DIR = os.path.join(RESULT_DIR, "code_shots")
REPORT_DIR = os.path.join(PROJECT, "report")

STUDENT_ID = "<学号>"          # 使用时可改成自己的学号
STUDENT_NAME = "<姓名>"        # 使用时可改成自己的姓名
TOPIC = "选题2：图像分类模型"
TITLE = "基于改进卷积神经网络的 CIFAR-10 图像分类模型"
REPORT_NAME = "实验报告.docx"

FONT_TITLE = "方正小标宋简体"      # 小二
FONT_BODY = "仿宋_GB2312"
FONT_HEAD = "黑体"

CLASS_CN = {"airplane": "飞机", "automobile": "汽车", "bird": "鸟", "cat": "猫", "deer": "鹿",
            "dog": "狗", "frog": "青蛙", "horse": "马", "ship": "轮船", "truck": "卡车"}


# =====================================================================
# 工具：把代码 / 终端输出渲染成“截图”
# =====================================================================
def _runs(text):
    out, cur, cur_ascii = [], "", None
    for ch in text:
        a = ord(ch) < 128
        if cur_ascii is None or a == cur_ascii:
            cur += ch
        else:
            out.append((cur, cur_ascii))
            cur = ch
        cur_ascii = a
    if cur:
        out.append((cur, cur_ascii))
    return out


def render_code(code_text, out_png, caption_title="code", size=13, pad=18, max_lines=60):
    lines = code_text.rstrip("\n").split("\n")[:max_lines]
    fa = ImageFont.truetype(r"C:\Windows\Fonts\consola.ttf", size)
    fc = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", int(size * 0.92))
    line_h = int(size * 1.6)
    probe = ImageDraw.Draw(Image.new("RGB", (10, 10)))

    def w_of(t):
        return sum(probe.textlength(s, font=fa if a else fc) for s, a in _runs(t))

    lineno_w = int(w_of("000") + 12)
    width = int(max([w_of(l) for l in lines] + [320]) + lineno_w + pad * 2 + 8)
    height = int(line_h * len(lines) + pad * 2 + 40)
    img = Image.new("RGB", (width, height), (30, 32, 36))
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, width, 26], fill=(44, 47, 52))
    d.ellipse([10, 8, 20, 18], fill=(255, 95, 86))
    d.ellipse([26, 8, 36, 18], fill=(255, 189, 46))
    d.ellipse([42, 8, 52, 18], fill=(39, 201, 63))
    d.text((60, 5), caption_title, font=fc, fill=(200, 200, 200))

    y = pad + 26
    for i, raw in enumerate(lines, 1):
        d.text((pad, y), f"{i:>3}", font=fa, fill=(110, 115, 125))
        x = pad + lineno_w
        st = raw.strip()
        if st.startswith("#") or st.startswith('"""'):
            color = (126, 168, 120)
        elif "'" in st and st.count("'") >= 2:
            color = (206, 145, 120)
        else:
            color = (212, 212, 212)
        for seg, a in _runs(raw):
            f = fa if a else fc
            d.text((x, y), seg, font=f, fill=color)
            x += probe.textlength(seg, font=f)
        y += line_h
    os.makedirs(os.path.dirname(out_png), exist_ok=True)
    img.save(out_png)
    return out_png


def render_console(lines, out_png, title, size=12, pad=14):
    fa = ImageFont.truetype(r"C:\Windows\Fonts\consola.ttf", size)
    fc = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", size)
    line_h = int(size * 1.55)
    probe = ImageDraw.Draw(Image.new("RGB", (10, 10)))
    max_w = max([sum(probe.textlength(s, font=fa if a else fc) for s, a in _runs(l)) for l in lines] + [400])
    width = int(max_w + pad * 2 + 6)
    height = int(line_h * len(lines) + pad * 2 + 32)
    img = Image.new("RGB", (width, height), (24, 24, 26))
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, width, 26], fill=(46, 46, 50))
    d.text((pad, 5), title, font=fc, fill=(190, 190, 190))
    y = pad + 28
    for line in lines:
        color = (214, 214, 214)
        if "保存最优模型" in line or "达标" in line:
            color = (120, 220, 150)
        elif line.strip().startswith("[Epoch"):
            color = (150, 200, 255)
        x = pad
        for seg, a in _runs(line):
            f = fa if a else fc
            d.text((x, y), seg, font=f, fill=color)
            x += probe.textlength(seg, font=f)
        y += line_h
    os.makedirs(os.path.dirname(out_png), exist_ok=True)
    img.save(out_png)
    return out_png


def read_lines(path):
    if not os.path.isfile(path):
        return []
    with open(path, encoding="utf-8", errors="ignore") as f:
        return [l.rstrip("\n") for l in f]


def extract_code(rel_path, start_marker, end_marker=None, max_lines=60):
    text = open(os.path.join(PROJECT, rel_path), encoding="utf-8").read()
    s = text.index(start_marker)
    e = text.index(end_marker, s) if end_marker else len(text)
    return "\n".join(text[s:e].rstrip().split("\n")[:max_lines])


# =====================================================================
# Word 文档工具
# =====================================================================
class Report:
    def __init__(self):
        self.doc = Document()
        self._page()
        self._styles()
        self.n_fig = {}
        self.n_tab = {}
        self.n_shot = {}

    def _page(self):
        for s in self.doc.sections:
            s.page_width, s.page_height = Cm(21.0), Cm(29.7)
            s.left_margin = s.right_margin = Cm(2.6)
            s.top_margin, s.bottom_margin = Cm(2.5), Cm(2.3)
            p = s.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._font(p.add_run("第 "), size=10.5)
            fld = OxmlElement("w:fldSimple")
            fld.set(qn("w:instr"), "PAGE")
            rr = OxmlElement("w:r"); tt = OxmlElement("w:t"); tt.text = "1"
            rr.append(tt); fld.append(rr); p._p.append(fld)
            self._font(p.add_run(" 页"), size=10.5)

    def _styles(self):
        n = self.doc.styles["Normal"]
        n.font.name = "Times New Roman"
        n.font.size = Pt(12)
        n.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        n.paragraph_format.line_spacing = 1.5
        n.paragraph_format.space_after = Pt(0)

    @staticmethod
    def _font(run, cn=FONT_BODY, en="Times New Roman", size=12, bold=False, color=None):
        run.font.name = en
        run.font.size = Pt(size)
        run.font.bold = bold
        rPr = run._element.get_or_add_rPr()
        rF = rPr.find(qn("w:rFonts"))
        if rF is None:
            rF = OxmlElement("w:rFonts"); rPr.append(rF)
        rF.set(qn("w:eastAsia"), cn)
        rF.set(qn("w:ascii"), en)
        rF.set(qn("w:hAnsi"), en)
        if color:
            run.font.color.rgb = RGBColor(*color)

    def title(self, text):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(8)
        self._font(p.add_run(text), cn=FONT_TITLE, size=18)

    def info(self, text):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        self._font(p.add_run(text), size=12)

    def h1(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
        self._font(p.add_run(text), cn=FONT_HEAD, size=15, bold=True)

    def h2(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
        self._font(p.add_run(text), cn=FONT_HEAD, size=13, bold=True)

    def h3(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)
        self._font(p.add_run(text), size=12.5, bold=True)

    def p(self, text, indent=True, size=12, bold=False, align=None):
        para = self.doc.add_paragraph()
        if indent:
            para.paragraph_format.first_line_indent = Pt(24)
        para.paragraph_format.space_after = Pt(3)
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for i, seg in enumerate(text.split("**")):
            if not seg:
                continue
            self._font(para.add_run(seg), size=size, bold=bold or i % 2 == 1)
        return para

    def bullet(self, items, size=12):
        for it in items:
            para = self.doc.add_paragraph(style="List Bullet")
            para.paragraph_format.line_spacing = 1.4
            para.paragraph_format.space_after = Pt(1)
            self._font(para.add_run(it), size=size)

    def caption(self, text):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(8)
        self._font(p.add_run(text), cn="楷体", size=10.5)

    def pic(self, path, cap, prefix="3", width=15.2):
        if not os.path.isfile(path):
            print(f"[跳过] 缺图 {os.path.relpath(path, PROJECT)}")
            return
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.add_run().add_picture(path, width=Cm(width))
        self.n_fig[prefix] = self.n_fig.get(prefix, 0) + 1
        self.caption(f"图 {prefix}-{self.n_fig[prefix]}  {cap}")

    def code_pic(self, code_text, cap, prefix="5", title="code.py", width=15.6, size=12):
        os.makedirs(CODE_SHOT_DIR, exist_ok=True)
        idx = self.n_fig.get(prefix, 0) + 1
        out = os.path.join(CODE_SHOT_DIR, f"code_{prefix}_{idx}.png")
        render_code(code_text, out, title, size=size)
        self.pic(out, cap, prefix=prefix, width=width)

    def console_pic(self, lines, cap, prefix="4", title="terminal", width=15.6, size=11):
        if not lines:
            return
        os.makedirs(SHOT_DIR, exist_ok=True)
        self.n_shot[prefix] = self.n_shot.get(prefix, 0) + 1
        out = os.path.join(SHOT_DIR, f"terminal_{prefix}_{self.n_shot[prefix]}.png")
        render_console(lines, out, title, size=size)
        self.pic(out, cap, prefix=prefix, width=width)

    def table(self, header, rows, cap, prefix="4", widths=None, size=10.5):
        tb = self.doc.add_table(rows=1, cols=len(header))
        tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, t in enumerate(header):
            c = tb.rows[0].cells[i]; c.text = ""
            pr = c.paragraphs[0]; pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pr.paragraph_format.line_spacing = 1.1
            self._font(pr.add_run(str(t)), cn=FONT_HEAD, size=size, bold=True)
            self._shade(c, "DCE6F1")
        for row in rows:
            cells = tb.add_row().cells
            for i, t in enumerate(row):
                cells[i].text = ""
                pr = cells[i].paragraphs[0]; pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pr.paragraph_format.line_spacing = 1.1
                self._font(pr.add_run(str(t)), size=size)
        if widths:
            for i, w in enumerate(widths):
                for r in tb.rows:
                    r.cells[i].width = Cm(w)
        self.n_tab[prefix] = self.n_tab.get(prefix, 0) + 1
        self.caption(f"表 {prefix}-{self.n_tab[prefix]}  {cap}")

    @staticmethod
    def _shade(cell, color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
        tcPr.append(shd)

    def brk(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.doc.save(path)
        print(f"[OK] Word 实验报告已生成：{path}")


# =====================================================================
# 生成报告
# =====================================================================
def load(name):
    with open(os.path.join(RESULT_DIR, name), encoding="utf-8") as f:
        return json.load(f)


def build():
    ds = load("dataset_stats.json")
    mi = load("metrics_improved.json")
    mb = load("metrics_baseline.json")
    hi = load("history_improved.json")
    hb = load("history_baseline.json")
    per = ds["train"]["per_class"]

    R = Report()

    # ---------------- 标题 / 姓名学号 ----------------
    R.title(TITLE)
    R.info(f"{STUDENT_NAME}　　学号：{STUDENT_ID}")

    # ---------------- 一、题目介绍 ----------------
    R.h1("一、题目介绍")
    R.p(f"本作品选择《考核方案》四个选题中的**{TOPIC}**，题目为《{TITLE}》。")
    R.p("考核要求：围绕课堂所学图像分类模型，自行选取不少于 1000 张、类别不少于 3 类的图像数据集完成训练与测试；"
        "对所用分类网络进行修改（增减层数、修改卷积核数量与尺寸等），不得直接使用课堂原始网络；"
        "用 Matplotlib 完成“第一行测试图片、第二行预测类别、第三行真实类别”的可视化并显示预测置信度。")
    R.p(f"完成情况：选用公开数据集 CIFAR-10（{ds['all_total']} 张、{ds['num_classes']} 类），"
        f"训练集 {ds['train']['total']} 张 / 验证集 {ds['val']['total']} 张 / 测试集 {ds['test']['total']} 张，"
        "三者完全独立、互不重叠，**测试集不参与任何训练与调参**。"
        f"改进网络训练 {hi['epochs']} 个 epoch（≥20 要求），测试集准确率 **{mi['test_acc']:.2%}**（≥85% 要求），"
        f"较课堂原始网络 {mb['test_acc']:.2%} 提升 {(mi['test_acc'] - mb['test_acc']) * 100:.2f} 个百分点，"
        f"分类头参数由 {mb['head_params']:,} 压缩到 {mi['head_params']:,}（轻量化）。")

    # ---------------- 二、所用工具介绍 ----------------
    R.h1("二、所用工具介绍")
    R.p("全部代码使用 Python 编写，主要依赖如下：")
    R.table(["工具 / 库", "版本", "在本作品中的作用"],
            [["Python", "3.10.21", "开发语言，所有脚本运行环境"],
             ["PyTorch", "2.14.0", "张量计算、自动求导、卷积网络搭建、训练与推理"],
             ["torchvision", "0.29.0", "ImageFolder 数据加载、transforms 数据增强与归一化"],
             ["NumPy", "2.2.6", "数值统计与数组运算"],
             ["Matplotlib", "3.10.9", "类别分布、训练曲线、混淆矩阵、三行预测、前后对比图"],
             ["Pillow", "12.3.0", "图片读取、保存与预览"],
             ["python-docx", "1.2.0", "自动生成本实验报告"]],
            "开发环境与依赖库", prefix="2", widths=[3.6, 2.8, 9.0], size=9.5)
    R.h2("（一）PyTorch")
    R.p("PyTorch 是本作品的核心框架，具体用到以下模块：")
    R.h3("1. 网络与损失函数")
    R.p("（1）`nn.Conv2d`：二维卷积，用于提取局部特征；（2）`nn.BatchNorm2d`：批归一化，稳定训练；"
        "（3）`nn.MaxPool2d` 与 `nn.AdaptiveAvgPool2d`：最大池化与全局平均池化；"
        "（4）`nn.Linear`、`nn.Dropout`：全连接层与随机失活；"
        "（5）`nn.CrossEntropyLoss(label_smoothing=0.1)`：带标签平滑的交叉熵损失。")
    R.h3("2. 数据与优化器")
    R.p("（1）`DataLoader`：批量读取、打乱、多进程加载（`persistent_workers=True` 避免 Windows 下每轮重启子进程）；"
        "（2）`torch.optim.SGD(momentum=0.9, nesterov=True)`：带动量的随机梯度下降；"
        "（3）`CosineAnnealingLR`：学习率余弦退火；（4）`torch.autocast('cpu', bfloat16)`：CPU 混合精度加速。")
    R.h2("（二）torchvision")
    R.p("（1）`datasets.ImageFolder`：按“根目录/类别/图片”结构自动加载并建立类别映射，训练与测试脚本共用，"
        "保证类别顺序一致；（2）`transforms`：`RandomCrop`、`RandomHorizontalFlip`、`ColorJitter`、"
        "`ToTensor`、`Normalize`、`RandomErasing` 等，训练集用随机增强、评估集只用归一化。")
    R.h2("（三）Matplotlib 与 python-docx")
    R.p("Matplotlib 负责全部图表（含本报告中的 11 张结果图）；python-docx 负责把真实指标与图表自动写入 Word，"
        "保证报告中的数据与代码运行结果完全一致。")

    # ---------------- 三、数据集准备 ----------------
    R.h1("三、数据集准备")
    R.h2("（一）数据集介绍与必要说明")
    R.p("数据来源：公开数据集 CIFAR-10（官网 https://www.cs.toronto.edu/~kriz/cifar.html），"
        "由 10 类真实彩色照片组成，每张 32×32 像素 RGB 图像，是图像分类最常用的基准数据集之一。"
        "遵照考核方案“不建议直接使用 MNIST”的要求，本作品未使用 MNIST。")
    R.p(f"数据规模与分布：共 {ds['all_total']} 张、{ds['num_classes']} 类，完全均衡——"
        f"训练集每类 {per['cat']} 张、验证集每类 {ds['val']['per_class']['cat']} 张、"
        f"测试集每类 {ds['test']['per_class']['cat']} 张（要求每类 ≥250 张、≥3 类，均已远超）。")
    R.table(["类别", "中文", "训练集", "验证集", "测试集", "合计"],
            [[n, cn, per[n], ds["val"]["per_class"][n], ds["test"]["per_class"][n],
              per[n] + ds["val"]["per_class"][n] + ds["test"]["per_class"][n]]
             for n, cn in zip(ds["class_names"], ds["class_names_cn"])]
            + [["合计", "10 类", ds["train"]["total"], ds["val"]["total"], ds["test"]["total"], ds["all_total"]]],
            "CIFAR-10 数据集类别分布（张）", prefix="3",
            widths=[3.0, 1.8, 2.4, 2.4, 2.4, 2.4], size=9.5)
    R.pic(os.path.join(RESULT_DIR, "dataset_distribution.png"), "三个子集里各类别的数量分布", prefix="3")
    R.pic(os.path.join(RESULT_DIR, "dataset_preview.png"),
          "训练集样本预览（每类随机 8 张，类内姿态、颜色、背景差异很大）", prefix="3")
    R.h2("（二）数据清洗与标注")
    R.p("CIFAR-10 为官方清洗并带标签的公开数据集，本作品的“数据制作”工作包括格式转换、校验与划分：")
    R.bullet(["（1）格式转换：把官方二进制数据导出为 PNG 图片，组织成 ImageFolder 目录结构；",
              "（2）完整性校验：导出时统计每类张数并逐张校验可读性，确认 60000 张全部可用；",
              f"（3）分层抽样划分：官方 50000 张训练图按类别打乱后 9:1 切分，得训练集 {ds['train']['total']} 张、"
              f"验证集 {ds['val']['total']} 张；官方 10000 张测试图作为独立测试集；三者零重叠，随机种子 seed=42 可复现；",
              "（4）数据增强：训练时随机裁剪、水平翻转、颜色抖动、随机擦除；验证/测试集只做归一化，保证评估公平。"])
    R.pic(os.path.join(RESULT_DIR, "augmentation_preview.png"),
          "同一张图片的 8 个随机增强变体（裁剪、翻转、颜色抖动）", prefix="3")
    R.h2("（三）图片展示与加载数据集代码截图")
    R.code_pic(extract_code("code/01_prepare_dataset.py", "def export_split", "def build_dataset", max_lines=24),
               "数据集导出代码（按类别写入 PNG，实现格式转换与校验）", prefix="3",
               title="01_prepare_dataset.py", size=12)
    R.code_pic(extract_code("code/01_prepare_dataset.py", "    rng = np.random.RandomState(SEED)",
                            "    print(\"[3/4]", max_lines=26),
               "9:1 分层抽样划分代码（train/val 来自官方训练图，test 独立，不参与训练）", prefix="3",
               title="01_prepare_dataset.py", size=12)
    R.code_pic(extract_code("code/03_train.py", "    train_tf = transforms.Compose([", "    ])", max_lines=14),
               "ImageFolder 加载数据集与数据增强代码", prefix="3", title="03_train.py", size=12)

    # ---------------- 四、完整训练和测试过程 ----------------
    R.h1("四、完整训练和测试过程")
    R.h2("（一）训练过程（训练参数）")
    R.p("训练脚本 `code/03_train.py`；原始网络与改进网络使用**完全相同**的训练配置，以保证控制变量对比：")
    R.table(["训练参数", "取值", "说明"],
            [["训练/验证/测试", f"{ds['train']['total']} / {ds['val']['total']} / {ds['test']['total']} 张",
              "三者独立，测试集不参与训练与调参"],
             ["训练轮数", f"改进 {hi['epochs']} / 基线 {hb['epochs']} epochs", "均满足 ≥20 轮要求"],
             ["批大小 batch_size", "128", "—"],
             ["优化器", "SGD + Nesterov（动量 0.9）", "小数据集泛化优于 Adam"],
             ["初始学习率", "0.1", "余弦退火逐步衰减"],
             ["权重衰减", "5e-4", "L2 正则抑制过拟合"],
             ["损失函数", "交叉熵 + 标签平滑 0.1", "避免过度自信"],
             ["数据增强", "随机裁剪/翻转/颜色抖动/随机擦除", "仅在训练集使用"],
             ["混合精度", "bf16（CPU AVX512）", "提速约 1.7 倍，权重仍 fp32"],
             ["最优权重选择", "按验证集准确率", "不使用测试集"]],
            "训练参数设置", prefix="4", widths=[3.8, 6.4, 5.4], size=9.5)
    imp_log = [l for l in read_lines(os.path.join(RESULT_DIR, "train_log_improved.txt")) if l.strip()]
    epochs = [l for l in imp_log if l.startswith("[Epoch")]
    R.p("训练过程输出如下（逐轮打印轮数、训练与验证指标、耗时与保存最优模型标记；"
        "训练阶段不接触测试集，完整记录见 `results/train_log_improved.txt`）：")
    R.console_pic(epochs[:8] + ["      ......（中间 epoch 省略）......"] + epochs[-4:]
                  + [l for l in imp_log if l.startswith("训练结束")],
                  "改进网络逐轮训练过程输出", prefix="4",
                  title="python 03_train.py --model improved --epochs 24", size=11)
    R.code_pic(extract_code("code/03_train.py", "    for epoch in range(1, args.epochs + 1):",
                            "    print(\"=\" * 84)", max_lines=32),
               "训练主循环代码（前向传播、损失计算、反向传播、验证集评估、保存最优权重）", prefix="4",
               title="03_train.py", size=12)
    R.pic(os.path.join(RESULT_DIR, "curve_improved.png"),
          "改进网络损失曲线与准确率曲线（训练集 vs 验证集，红虚线为 85% 要求）", prefix="4")
    R.p(f"曲线分析：损失单调下降并趋于平缓，准确率同步上升，说明正常收敛；训练准确率 {hi['train_acc'][-1]:.2%}、"
        f"验证准确率 {hi['val_acc'][-1]:.2%}，相差 {abs(hi['train_acc'][-1] - hi['val_acc'][-1]) * 100:.1f} 个百分点，"
        "没有“训练涨、验证掉”的剪刀差，配合数据增强、标签平滑、Dropout 与权重衰减，未出现明显过拟合。"
        f"最优权重来自第 {hi['best_epoch']} 个 epoch（验证集 {hi['best_val_acc']:.2%}）。")

    R.h2("（二）测试过程与代码截图")
    R.p("测试脚本 `code/04_test_and_visualize.py`：加载最优权重，在**从未参与训练**的 10000 张测试集上逐张预测，"
        "统计总体准确率与每类精确率/召回率/F1，并绘制混淆矩阵与三行预测可视化。测试过程输出：")
    test_log = [l for l in read_lines(os.path.join(RESULT_DIR, "console_test_improved.txt")) if l.strip()]
    R.console_pic(test_log[:24], "测试过程输出（权重来源、测试集规模、总体准确率、每类指标）",
                  prefix="4", title="python 04_test_and_visualize.py --model improved", size=11)
    R.code_pic(extract_code("code/04_test_and_visualize.py", "def evaluate(", "def per_class_metrics", max_lines=24),
               "测试集评估代码（逐批前向、Softmax、混淆矩阵与准确率统计）", prefix="4",
               title="04_test_and_visualize.py", size=12)
    R.p(f"测试结果：测试集 {mi['test_total']} 张，总体准确率 **{mi['test_acc']:.2%}**，"
        f"宏平均 F1 **{mi['macro_f1']:.4f}**，平均预测置信度 {mi['mean_confidence']:.2%}。每类指标：")
    R.table(["类别", "中文", "测试张数", "预测正确", "精确率", "召回率", "F1"],
            [[r["class"], r["class_cn"], r["support"], r["correct"],
              f"{r['precision']:.2%}", f"{r['recall']:.2%}", f"{r['f1']:.4f}"] for r in mi["per_class"]]
            + [["宏平均", "—", mi["test_total"], "—", "—", "—", f"{mi['macro_f1']:.4f}"]],
            f"改进网络在独立测试集上的每类指标（总体 {mi['test_acc']:.2%}）", prefix="4",
            widths=[2.6, 1.6, 2.2, 2.2, 2.4, 2.4, 2.0], size=9.5)

    # ---------------- 五、方法说明 ----------------
    R.h1("五、方法说明：模型介绍与修改改进")
    R.h2("（一）模型介绍（课堂原始网络）")
    R.p("课堂原始网络（BaselineCNN）：3 个「Conv3×3 + BN + ReLU + MaxPool」堆叠，通道 16/32/64，"
        "后接 `Flatten → Linear(1024→128) → ReLU → Dropout(0.3) → Linear(128→10)`。"
        f"总参数 {mb['params']:,}，其中分类头 {mb['head_params']:,}（占 {mb['head_params'] / mb['params']:.1%}）。")
    R.code_pic(extract_code("code/02_models.py", "class BaselineCNN", "class ImprovedCNN", max_lines=26),
               "课堂原始网络 BaselineCNN 代码", prefix="5", title="02_models.py", size=12)
    R.table(["层（按顺序）", "输出尺寸", "说明"],
            [["Conv2d(3→16,3×3,p=1)+BN+ReLU + MaxPool", "16×16×16", "第 1 个卷积块"],
             ["Conv2d(16→32,3×3,p=1)+BN+ReLU + MaxPool", "32×8×8", "第 2 个卷积块"],
             ["Conv2d(32→64,3×3,p=1)+BN+ReLU + MaxPool", "64×4×4", "第 3 个卷积块"],
             ["Flatten + Linear(1024→128)+ReLU+Dropout(0.3)", "128", "全连接层"],
             ["Linear(128→10)", "10", "输出 10 类得分"]],
            "BaselineCNN 结构（修改前）", prefix="5", widths=[8.0, 3.2, 3.8], size=9.5)

    R.h2("（二）模型改进展示（层数、卷积核数量与尺寸的修改）")
    R.p("针对基线“卷积层过浅、分类头参数占比过高、特征图压缩过快”的问题，设计 ImprovedCNN，共 6 处修改：")
    R.table(["序号", "修改内容", "修改目的"],
            [["①", "卷积层 3 层 → 8 层", "加深网络，实现“边缘→纹理→部件→物体”的层次化特征"],
             ["②", "通道数 16/32/64 → 32/64/128/256", "特征图更宽，保留更多细节"],
             ["③", "每个分辨率阶段堆叠 2 个 3×3 卷积", "两个 3×3 的感受野≈一个 5×5，但参数更少、多一层非线性"],
             ["④", "末端新增 1×1 卷积（128→256）", "仅做通道方向线性组合，参数极少，用于升维与增加非线性"],
             ["⑤", "Flatten+大全连接 → 全局平均池化 + Linear(256→10)",
              f"分类头参数 {mb['head_params']:,} → {mi['head_params']:,}（-98%），轻量化且抗过拟合"],
             ["⑥", "训练策略：标签平滑、余弦退火、更强增强", "提升泛化（与结构修改分开说明）"]],
            "网络修改点（对应考核要求：增减层数、修改卷积核数量和尺寸）", prefix="5",
            widths=[1.2, 6.8, 7.0], size=9.5)
    R.code_pic(extract_code("code/02_models.py", "class ImprovedCNN", "        return self.classifier", max_lines=40),
               "改进网络 ImprovedCNN 代码（8 层卷积 + 1×1 卷积 + 全局平均池化）", prefix="5",
               title="02_models.py", size=11)
    R.table(["层（按顺序）", "输出尺寸", "说明"],
            [["Conv2d(3→32,3×3)+BN+ReLU", "32×32×32", "第 1 层：颜色与边缘"],
             ["Conv2d(32→32,3×3)×2 + BN+ReLU + MaxPool", "32×16×16", "第 2~3 层 + 池化"],
             ["Conv2d(32→64,3×3)×2 + BN+ReLU + MaxPool", "64×8×8", "第 4~5 层 + 池化"],
             ["Conv2d(64→128,3×3)×2 + BN+ReLU + MaxPool", "128×4×4", "第 6~7 层 + 池化"],
             ["Conv2d(128→256,1×1)+BN+ReLU", "256×4×4", "第 8 层：1×1 卷积升维"],
             ["AdaptiveAvgPool2d(1) → Dropout(0.4) → Linear(256→10)", "10", "轻量分类头（2,570 参数）"]],
            "ImprovedCNN 结构（修改后）", prefix="5", widths=[8.0, 3.2, 3.8], size=9.5)
    R.p("**数据传递过程与模型原理：** 输入 3×32×32 图片先由第 1 个 3×3 卷积提取颜色与边缘；每经过一次池化，"
        "特征图边长减半、通道翻倍，网络关注的层次从“线条”逐步升级为“纹理”“部件”，最后成为“整只动物/整辆车”的语义特征。"
        "末端 1×1 卷积在通道方向做线性组合完成升维；全局平均池化把每个通道压成一个标量，得到 256 维特征向量；"
        "Dropout 随机丢弃部分特征防止死记；最后由 256→10 线性层输出 10 个类别的得分，"
        "经 Softmax 变成概率分布，概率最大的类别即预测结果，该概率即“预测置信度”。"
        "这样设计的作用是：以较少参数获得更深的特征层次（精度提升），同时把参数集中在真正做特征提取的卷积层、"
        "把分类头压到极轻（轻量化），最终实现在 32×32 低分辨率照片上稳定区分 10 个相似类别的效果。")
    R.p(f"**精度与轻量化改进思路（高分要点）：** 在完全相同训练配置下，改进网络测试准确率由 "
        f"{mb['test_acc']:.2%} 提升至 **{mi['test_acc']:.2%}**（+{(mi['test_acc'] - mb['test_acc']) * 100:.2f} 个百分点），"
        f"验证集由 {mb['val_acc']:.2%} 提升至 **{mi['val_acc']:.2%}**（+{(mi['val_acc'] - mb['val_acc']) * 100:.2f} 个百分点）；"
        f"总参数仅 {mi['params']:,}（约 0.33M），分类头从 {mb['head_params']:,} 降至 {mi['head_params']:,}（-98%），"
        "权重文件 1.3 MB，CPU 单轮训练约 1 分钟、单张预测毫秒级，精度与轻量化同时改进。")

    # ---------------- 六、效果展示 ----------------
    R.h1("六、效果展示")
    R.h2("（一）效果代码展示")
    R.p("预测可视化代码（严格按要求：第一行测试图片、第二行预测类别与置信度、第三行真实类别，"
        "边框绿色表示预测正确、红色表示错误）：")
    R.code_pic(extract_code("code/04_test_and_visualize.py", "def plot_prediction_grid", "def test_model",
                            max_lines=42),
               "三行预测可视化代码", prefix="6", title="04_test_and_visualize.py", size=11)
    R.h2("（二）可视化图片展示与对比截图")
    R.pic(os.path.join(RESULT_DIR, "confusion_matrix_improved.png"),
          "改进网络在测试集上的混淆矩阵（左：数量；右：按真实类别归一化）", prefix="6")
    R.pic(os.path.join(RESULT_DIR, "prediction_grid_improved.png"),
          "预测结果三行可视化：第一行测试图片、第二行预测类别与置信度、第三行真实类别", prefix="6")
    R.pic(os.path.join(RESULT_DIR, "comparison.png"),
          "网络修改前后测试结果对比截图（左：测试准确率；中：验证准确率曲线；右：每类准确率）", prefix="6")
    R.table(["指标", "修改前 BaselineCNN", "修改后 ImprovedCNN", "变化"],
            [["卷积层数", "3", "8", "+5 层"],
             ["卷积核尺寸/数量", "3×3，3 个卷积层", "3×3 共 7 层 + 1×1 共 1 层", "加深加宽 + 1×1 卷积"],
             ["总参数量", f"{mb['params']:,}", f"{mi['params']:,}", f"{mi['params'] / mb['params']:.2f} 倍"],
             ["分类头参数", f"{mb['head_params']:,}", f"{mi['head_params']:,}",
              f"减少 {1 - mi['head_params'] / mb['head_params']:.1%}（轻量化）"],
             ["训练轮数", f"{hb['epochs']}", f"{hi['epochs']}", "均 ≥20 轮"],
             ["测试集准确率", f"{mb['test_acc']:.2%}", f"{mi['test_acc']:.2%}",
              f"+{(mi['test_acc'] - mb['test_acc']) * 100:.2f} 个百分点"],
             ["验证集准确率", f"{mb['val_acc']:.2%}", f"{mi['val_acc']:.2%}",
              f"+{(mi['val_acc'] - mb['val_acc']) * 100:.2f} 个百分点"],
             ["宏平均 F1", f"{mb['macro_f1']:.4f}", f"{mi['macro_f1']:.4f}",
              f"+{mi['macro_f1'] - mb['macro_f1']:.4f}"]],
            "网络修改前后测试结果对比", prefix="6", widths=[3.8, 4.2, 4.2, 3.6], size=9.5)
    R.h2("（三）效果分析")
    worst = min(mi["per_class"], key=lambda r: r["recall"])
    best = max(mi["per_class"], key=lambda r: r["recall"])
    cm, names = mi["confusion_matrix"], mi["class_names"]
    pairs = sorted([(cm[i][j], names[i], names[j]) for i in range(len(names))
                    for j in range(len(names)) if i != j and cm[i][j] > 0], reverse=True)[:3]
    top3 = "、".join(f"{CLASS_CN[a]}→{CLASS_CN[b]} {n} 张" for n, a, b in pairs)
    R.p(f"（1）整体：独立测试集准确率 {mi['test_acc']:.2%}，宏平均 F1 {mi['macro_f1']:.4f}；"
        f"识别最好的是 {best['class']}（{best['class_cn']}，{best['recall']:.2%}），"
        f"最差的是 {worst['class']}（{worst['class_cn']}，{worst['recall']:.2%}）。")
    R.p(f"（2）错误分布：主要原因集中在 {top3}。这些类别在 32×32 分辨率下轮廓高度相似"
        "（猫与狗、鹿与马、汽车与卡车），属数据集固有难点。")
    R.p(f"（3）收敛与过拟合：训练 {hi['train_acc'][-1]:.2%} 与验证 {hi['val_acc'][-1]:.2%} 同步上升、无剪刀差，"
        "损失曲线平滑下降，说明收敛正常、无明显过拟合。")
    R.p("（4）轻量化：分类头参数减少 98%，权重仅 1.3 MB，CPU 单轮约 1 分钟、单张预测毫秒级，具备边缘部署潜力。")
    R.h2("（四）后续可改进内容")
    R.bullet(["（1）引入残差连接（ResNet 结构），可继续加深网络而不退化，预计再提升 1~3 个百分点；",
              "（2）使用 Mixup / CutMix 等更强增强，进一步抑制过拟合；",
              "（3）模型集成或知识蒸馏，用大模型能力提升小模型精度；",
              "（4）用 torch.quantization 量化、导出 ONNX，压缩模型以便部署到无人机等边缘设备；",
              "（5）与目标检测结合，先定位再细分，提升复杂场景实用性。"])

    # ---------------- 七、个人总结心得 ----------------
    R.h1("七、个人总结心得")
    R.p("这次期末作品让我完整走了一遍“数据 → 模型 → 训练 → 评估 → 报告”的全流程，收获主要有三点。"
        "一是真正理解了几条纪律的意义：训练集、验证集、测试集必须相互独立，测试集绝不能参与训练与调参——"
        "以前只觉得是形式，自己动手划分数据、看曲线时才明白，这条纪律决定了成绩是否可信。"
        "二是对网络结构有了“手感”：把课堂的 3 层小网络改成 8 层、加宽通道、加入 1×1 卷积、"
        "用全局平均池化替掉大全连接后，分类头参数从 13 万降到 2 千多，准确率反而更高，"
        "这让我体会到“结构设计要有目的，不是盲目堆参数”。"
        "三是训练技巧与工程细节的影响远超预期：数据增强、标签平滑、余弦退火、Dropout 与权重衰减，"
        "每一项都能在曲线上看出差别；而 Windows 下 DataLoader 每轮重启子进程会拖慢训练、"
        "CPU 上开启 bf16 混合精度能提速近两倍、控制台中文要用 UTF-8 写日志——"
        "这些课堂上一句带过的细节，真正做起来最卡人。")
    R.p("不足也很明显：受 CPU 算力限制只训练了 24 轮，没有做更多消融实验（例如单独去掉某项增强看影响）；"
        "模型仍是自己设计的轻量网络，未尝试残差、注意力等结构；可视化还可以更丰富（每类 PR 曲线、错误样本专题分析）。"
        "给老师的建议是：希望增加一次 GPU 环境的集体实验，让大家体验更大模型与更长训练；"
        "也可以安排“同学互评代码”的环节，看别人的代码往往比自己闷头写收获更大。"
        "总的来说，这门课让我从“会调用几行 API”变成“能独立完成一个小型 AI 项目并讲清原理”，"
        "把知识变成可验收的作品，是我最大的收获。")

    R.save(os.path.join(REPORT_DIR, REPORT_NAME))


if __name__ == "__main__":
    build()
